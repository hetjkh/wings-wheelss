from pathlib import Path
import sys, re, io, json, math, html
BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE.parent/'peb-libs'))
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import pymupdf as fitz
from PIL import Image
from report_content import PAGES

OUT=Path(r'D:\my-wings\output\peb-training-report'); FIG=OUT/'figures'
NAVY='#183247';TEAL='#087e8b';GREY='#637785';LIGHT='#eef4f6';ORANGE='#a1542c'
for name,file in [('Calibri','calibri.ttf'),('Calibri-Bold','calibrib.ttf'),('Calibri-Italic','calibrii.ttf')]:
    pdfmetrics.registerFont(TTFont(name,'C:/Windows/Fonts/'+file))
pdfmetrics.registerFontFamily('Calibri',normal='Calibri',bold='Calibri-Bold',italic='Calibri-Italic',boldItalic='Calibri-Bold')
PAGEW,PAGEH=A4; MARGIN=43; WIDTH=PAGEW-2*MARGIN

def plain(s):return html.unescape(re.sub('<[^>]+>','',str(s)))
def markup(s):return str(s).replace('\n','<br/>')
def style(size=10.6,leading=14.1,bold=False,color=NAVY,**kw):
    return ParagraphStyle('s',fontName='Calibri-Bold' if bold else 'Calibri',fontSize=size,leading=leading,textColor=HexColor(color),spaceAfter=kw.pop('spaceAfter',7),**kw)

def footer(c,doc,num):
    c.saveState();c.setStrokeColor(HexColor(TEAL));c.setLineWidth(1.1);c.line(MARGIN,PAGEH-36,PAGEW-MARGIN,PAGEH-36)
    if num>1:
        c.setFont('Calibri',8);c.setFillColor(HexColor(GREY));c.drawString(MARGIN,PAGEH-27,'STRUCTURAL ENGINEERING TRAINING  |  PEB CASE STUDY')
    c.setStrokeColor(HexColor('#cbd9df'));c.setLineWidth(.5);c.line(MARGIN,40,PAGEW-MARGIN,40)
    c.setFont('Calibri',8);c.setFillColor(HexColor(GREY));c.drawString(MARGIN,27,'DEMO REPORT  •  Educational study / not for construction')
    c.drawRightString(PAGEW-MARGIN,27,f'{num:02d} / 20');c.restoreState()

def image_flow(name,height):
    path=FIG/(name+'.png');im=Image.open(path);ratio=im.width/im.height
    w=min(WIDTH,height*ratio);hh=w/ratio
    return RLImage(str(path),width=w,height=hh,hAlign='CENTER')

def build_flows(pg,num,scale=1):
    body=style(size=10.6*scale,leading=14.0*scale)
    small=style(size=8.8*scale,leading=11.3*scale,color=GREY)
    heading=style(size=12.5*scale,leading=15*scale,bold=True,spaceBefore=5)
    flows=[]
    if num==1:
        flows += [Spacer(1,36),Paragraph('FOUR-WEEK STUDY  /  DEMO EDITION',style(10,14,bold=True,color=TEAL)),Spacer(1,10),Paragraph('STRUCTURAL ENGINEERING<br/>TRAINING REPORT',style(29,33,bold=True)),Paragraph('Analysis, design and detailing of a<br/>pre-engineered steel building',style(18,23)),Spacer(1,6),image_flow('model',200)]
        rows=[['Project','Proposed PEB industrial building, Chennai'],['Student','Demo Student  |  Roll No. DEMO-001'],['Institution','Demo College of Engineering'],['Training organization','Demo Structural Consultants'],['Supervisor','Demo Training Supervisor'],['Training period','Four weeks — dates to be entered'],['Report issue','21 October 2026  |  Revision D0']]
        t=Table([[Paragraph(a,style(10,13,bold=True)),Paragraph(b,style(10.5,13))] for a,b in rows],colWidths=[127,WIDTH-127]);t.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('ROWBACKGROUNDS',(0,0),(-1,-1),[HexColor(LIGHT),white])]))
        flows += [t,Spacer(1,15),Paragraph('36 m × 72.2 m  •  10 m eave level  •  Two 18 m roof spans',style(11,14,bold=True,color=TEAL)),Paragraph('Prepared from the supplied training PDF and STAAD.Pro analysis output. Personal and institutional details are demonstration entries. Source findings and unresolved design items are retained for technical accuracy.',small)]
        return flows
    flows += [Paragraph(pg['kicker'],style(9.4,12,bold=True,color=TEAL)),Spacer(1,3),Paragraph(markup(pg['title']),style(21,25,bold=True,spaceAfter=12))]
    for b in pg['blocks']:
        k=b['kind']
        if k=='p':flows.append(Paragraph(markup(b['text']),body))
        elif k=='h':flows.append(Paragraph(markup(b['text']),heading))
        elif k=='source':flows += [Spacer(1,3),Paragraph(markup('Source basis: '+b['text']),style(8.1*scale,10.5*scale,color=GREY))]
        elif k in ('eq','note'):
            pp=Paragraph(markup(b['text']),style((10.1 if k=='eq' else 9.5)*scale,(13.5 if k=='eq' else 12.5)*scale,color=NAVY if k=='eq' else ORANGE))
            tt=Table([[pp]],colWidths=[WIDTH]);tt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor(LIGHT if k=='eq' else '#faf1e9')),('BOX',(0,0),(-1,-1),.4,HexColor('#d6e3e8' if k=='eq' else '#e8d4c4')),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),5)]));flows += [tt,Spacer(1,7)]
        elif k=='table':
            widths=b['widths'] or [1/len(b['headers'])]*len(b['headers'])
            hs=style(9.1*scale,11.4*scale,bold=True,color='#ffffff',spaceAfter=0)
            ts=style(9.25*scale,11.8*scale,spaceAfter=0)
            data=[[Paragraph(markup(v),hs) for v in b['headers']]]+[[Paragraph(markup(v),ts) for v in row] for row in b['rows']]
            tt=Table(data,colWidths=[WIDTH*x/sum(widths) for x in widths],repeatRows=1,hAlign='LEFT')
            tt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor(NAVY)),('ROWBACKGROUNDS',(0,1),(-1,-1),[HexColor(LIGHT),white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),4.4*scale),('BOTTOMPADDING',(0,0),(-1,-1),4.4*scale),('LINEBELOW',(0,0),(-1,-1),.35,HexColor('#d5e1e6'))]))
            flows += [tt,Spacer(1,8)]
        elif k=='fig':
            flows += [image_flow(b['name'],b['height']*scale),Spacer(1,4),Paragraph(markup(b['caption']),small),Spacer(1,4)]
        elif k=='reference':
            label=html.escape(b['label'])
            if b['url']:label=f'<link href="{html.escape(b["url"],quote=True)}" color="{TEAL}">{label}</link>'
            flows += [Paragraph('<b>'+label+'</b> '+html.escape(b['text']),style(9.3*scale,12.1*scale,spaceAfter=7))]
    return flows

merged=fitz.open(); scales=[]
for i,pg in enumerate(PAGES,1):
    scale=1
    for attempt in range(10):
        buffer=io.BytesIO();doc=SimpleDocTemplate(buffer,pagesize=A4,leftMargin=MARGIN,rightMargin=MARGIN,topMargin=50,bottomMargin=51,title='PEB Structural Engineering Training Report')
        doc.build(build_flows(pg,i,scale),onFirstPage=lambda c,d:footer(c,d,i),onLaterPages=lambda c,d:footer(c,d,i))
        test=fitz.open(stream=buffer.getvalue(),filetype='pdf')
        if len(test)==1:break
        scale-=.025
    if len(test)!=1:raise RuntimeError(f'Page {i} does not fit')
    scales.append(scale);merged.insert_pdf(test);print(f'Page {i:02}: scale {scale:.3f} / {len(test[0].get_text().split())} words')
merged.set_metadata({'title':'Structural Engineering Training Report — PEB Design','author':'Demo Student','subject':'Four-week academic study based on supplied project records','keywords':'PEB, STAAD.Pro, steel design, training report, demo'})
merged.set_toc([[1,f'{i+1:02}  {p["title"].replace(chr(10)," ")}',i+1] for i,p in enumerate(PAGES)])
pdfpath=OUT/'PEB_Structural_Engineering_Training_Report.pdf';merged.save(pdfpath)

# Editable Word version, with the same 20 logical pages and smaller body spacing.
doc=Document(); sec=doc.sections[0];sec.page_width=Pt(PAGEW);sec.page_height=Pt(PAGEH)
sec.top_margin=Pt(48);sec.bottom_margin=Pt(45);sec.left_margin=Pt(MARGIN);sec.right_margin=Pt(MARGIN);sec.header_distance=Pt(22);sec.footer_distance=Pt(22)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.font.color.rgb=RGBColor.from_string(NAVY[1:]);normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.03
for nm,size in [('Title',27),('Heading 1',20),('Heading 2',12)]:
    st=doc.styles[nm];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(NAVY[1:]);st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(4)
hp=sec.header.paragraphs[0];hp.text='STRUCTURAL ENGINEERING TRAINING  |  PEB CASE STUDY';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string(GREY[1:])
fp=sec.footer.paragraphs[0];fp.text='DEMO REPORT • Educational study / not for construction                                      ';fp.runs[0].font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)

def rich(par,text,size=None):
    # Preserve simple inline emphasis, while keeping every paragraph editable.
    parts=re.split(r'(<b>|</b>|<i>|</i>)',str(text));bold=False;ital=False
    for s in parts:
        if s=='<b>':bold=True
        elif s=='</b>':bold=False
        elif s=='<i>':ital=True
        elif s=='</i>':ital=False
        else:
            r=par.add_run(plain(s));r.bold=bold;r.italic=ital
            if size:r.font.size=Pt(size)
    return par
def dp(text,sty=None,size=None):return rich(doc.add_paragraph(style=sty),text,size)
def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),fill.lstrip('#'));pr.append(el)
def dt(headers,rows,widths=None,scale=1):
    widths=widths or [1/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for j,(cell,head) in enumerate(zip(t.rows[0].cells,headers)):
        shade(cell,NAVY);pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);rich(pp,head,8.8*scale)
        for rr in pp.runs:rr.bold=True;rr.font.color.rgb=RGBColor(255,255,255)
    for idx,row in enumerate(rows):
        cells=t.add_row().cells
        for c,txt in zip(cells,row):
            if idx%2==0:shade(c,LIGHT)
            pp=c.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.0;rich(pp,txt,8.8*scale);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in t.rows:
        for j,c in enumerate(row.cells):c.width=Pt(WIDTH*widths[j]/sum(widths))
        trpr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trpr.append(el)
    for col,w in zip(t.columns,widths):col.width=Pt(WIDTH*w/sum(widths))
    dp('',size=1).paragraph_format.space_after=Pt(2)
def di(name,height):
    im=Image.open(FIG/(name+'.png'));w=min(WIDTH,height*im.width/im.height)
    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_after=Pt(3);pp.add_run().add_picture(str(FIG/(name+'.png')),width=Pt(w))
def link(par,label,url):
    rel=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rel);run=OxmlElement('w:r');props=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),TEAL[1:]);props.append(col);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');props.append(sz);run.append(props);tt=OxmlElement('w:t');tt.text=label;run.append(tt);el.append(run);par._p.append(el)
for idx,pg in enumerate(PAGES):
    if idx:doc.add_page_break()
    sc=scales[idx]
    if idx==0:
        dp('FOUR-WEEK STUDY / DEMO EDITION',size=10)
        dp('STRUCTURAL ENGINEERING\nTRAINING REPORT','Title')
        dp('Analysis, design and detailing of a\npre-engineered steel building',size=18)
        di('model',195)
        dt(['Report particulars','Demo / source details'],[['Project','Proposed PEB industrial building, Chennai'],['Student','Demo Student | Roll No. DEMO-001'],['Institution','Demo College of Engineering'],['Training organization','Demo Structural Consultants'],['Supervisor','Demo Training Supervisor'],['Training period','Four weeks — dates to be entered'],['Report issue','21 October 2026 | Revision D0']],[.28,.72])
        dp('36 m × 72.2 m • 10 m eave level • Two 18 m roof spans',size=12)
        dp('Prepared from the supplied training PDF and STAAD.Pro analysis output. Personal and institutional details are demonstration entries. Source findings and unresolved design items are retained for technical accuracy.',size=9)
        continue
    pp=dp(pg['kicker'],size=9);pp.runs[0].font.color.rgb=RGBColor.from_string(TEAL[1:]);dp(pg['title'],'Heading 1')
    for b in pg['blocks']:
        k=b['kind']
        if k=='p':dp(b['text'],size=10*sc)
        elif k=='h':dp(b['text'],'Heading 2')
        elif k=='source':dp('Source basis: '+b['text'],size=7.8*sc)
        elif k in ['eq','note']:
            pp=dp(b['text'],size=(9.7 if k=='eq' else 9.2)*sc);pp.paragraph_format.left_indent=Pt(6);pp.paragraph_format.right_indent=Pt(6)
            pr=pp._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF4F6' if k=='eq' else 'FAF1E9');pr.append(sh)
        elif k=='table':dt(b['headers'],b['rows'],b['widths'],sc)
        elif k=='fig':di(b['name'],b['height']*sc);dp(b['caption'],size=8.4*sc)
        elif k=='reference':
            pp=doc.add_paragraph()
            if b['url']:link(pp,b['label'],b['url'])
            else:rich(pp,'<b>'+b['label']+'</b>',9*sc)
            rich(pp,' '+b['text'],9*sc)
doc.core_properties.title='Structural Engineering Training Report — PEB Design';doc.core_properties.author='Demo Student';doc.core_properties.subject='Four-week study based on supplied project records'
doc.save(OUT/'PEB_Structural_Engineering_Training_Report.docx')
(OUT/'report-content.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2),encoding='utf-8')
(BASE/'layout-scales.json').write_text(json.dumps(scales))

# Render pages for verification and gather output statistics.
review=BASE/'review';review.mkdir(exist_ok=True)
for i,pdfpage in enumerate(merged):
    pdfpage.get_pixmap(matrix=fitz.Matrix(1.1,1.1)).save(review/f'page-{i+1:02}.png')
from PIL import ImageDraw
for batch in range(4):
    montage=Image.new('RGB',(5*310,460),'#dce5e8');dr=ImageDraw.Draw(montage)
    for j in range(5):
        n=batch*5+j+1;im=Image.open(review/f'page-{n:02}.png');im.thumbnail((300,425));montage.paste(im,(j*310+5,22));dr.text((j*310+8,4),f'PAGE {n}',fill='black')
    montage.save(review/f'contact-{batch+1}.jpg')
print('FINAL PDF PAGES:',len(merged))
print('Total extracted words:',sum(len(x.get_text().split()) for x in merged))
print('Purlin I example:',5*1.5127*8100**4/(384*205000*33.75))
print('Output:',str(OUT))
