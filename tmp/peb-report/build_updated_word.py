from pathlib import Path
import sys,json,re,html,zipfile,shutil
BASE=Path(__file__).resolve().parent
sys.path.insert(0,str(BASE.parent/'peb-libs'))
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

OUT=Path(r'D:\my-wings\output\peb-training-report');FIG=OUT/'figures'
pages=json.loads((OUT/'report-content-updated.json').read_text(encoding='utf-8'))
scales=json.loads((OUT/'validation-updated.json').read_text())['page_scales']
doc=Document();sec=doc.sections[0]
sec.page_width=Pt(595.276);sec.page_height=Pt(841.89)
sec.left_margin=sec.right_margin=Pt(53);sec.top_margin=Pt(50);sec.bottom_margin=Pt(43);sec.footer_distance=Pt(21)
CW=489.276
for name,size in [('Normal',10.8),('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11.4)]:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=Pt(13.7 if name=='Normal' else size+3)
    if name.startswith('Heading'):s.font.bold=True
doc.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY

def xml(tag,**attrs):
    e=OxmlElement('w:'+tag)
    for k,v in attrs.items():e.set(qn('w:'+k),str(v))
    return e
def rich(p,text,size=None):
    bold=False;italic=False
    for s in re.split(r'(<b>|</b>|<i>|</i>)',str(text)):
        if s=='<b>':bold=True
        elif s=='</b>':bold=False
        elif s=='<i>':italic=True
        elif s=='</i>':italic=False
        else:
            r=p.add_run(html.unescape(re.sub('<[^>]+>','',s)));r.bold=bold;r.italic=italic
            if size:r.font.size=Pt(size)
    return p
def para(text='',size=None,style=None):return rich(doc.add_paragraph(style=style),text,size)
def field(p,code,cached):
    r=OxmlElement('w:r');r.append(xml('fldChar',fldCharType='begin'));p._p.append(r)
    r=OxmlElement('w:r');i=OxmlElement('w:instrText');i.set(qn('xml:space'),'preserve');i.text=' '+code+' ';r.append(i);p._p.append(r)
    r=OxmlElement('w:r');r.append(xml('fldChar',fldCharType='separate'));p._p.append(r)
    p.add_run(str(cached))
    r=OxmlElement('w:r');r.append(xml('fldChar',fldCharType='end'));p._p.append(r)
def bookmark(p,n):
    p._p.insert(0,xml('bookmarkStart',id=n,name=f'report_page_{n}'));p._p.append(xml('bookmarkEnd',id=n))
def band(title,n):
    p=para(title,style='Heading 1');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before=n>1;p.paragraph_format.space_after=Pt(13)
    p.paragraph_format.space_before=Pt(7);p.paragraph_format.line_spacing=Pt(18)
    b=xml('pBdr')
    for pos in ['top','bottom']:b.append(xml(pos,val='single',sz=5,space=8,color='4668AC'))
    p._p.get_or_add_pPr().append(b);bookmark(p,n)
    return p
def picture(name,height):
    im=Image.open(FIG/(name+'.png'));w=min(CW,height*im.width/im.height)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(FIG/(name+'.png')),width=Pt(w))
def table(headers,rows,widths=None,sc=1):
    widths=widths or [1/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for vals,row in [(headers,t.rows[0])]+[(r,t.add_row()) for r in rows]:
        ishead=row._tr is t.rows[0]._tr
        for j,(c,txt) in enumerate(zip(row.cells,vals)):
            c.width=Pt(CW*widths[j]/sum(widths));c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            pr=c._tc.get_or_add_tcPr();marg=xml('tcMar')
            for side,value in [('top',45),('bottom',45),('left',80),('right',80)]:marg.append(xml(side,w=value,type='dxa'))
            pr.append(marg)
            if ishead:pr.append(xml('shd',fill='DCE3F1'))
            pp=c.paragraphs[0];pp.alignment=WD_ALIGN_PARAGRAPH.LEFT;pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(11.5*sc)
            rich(pp,('<b>'+txt+'</b>') if ishead else txt,9.3*sc)
        row._tr.get_or_add_trPr().append(xml('cantSplit'))
    borders=xml('tblBorders')
    for side in ['top','bottom','left','right','insideH','insideV']:borders.append(xml(side,val='single',sz=3,color='B4B4B4'))
    t._tbl.tblPr.append(borders)
    for c,w in zip(t.columns,widths):c.width=Pt(CW*w/sum(widths))
    p=para('',size=1);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=Pt(1)

fp=sec.footer.paragraphs[0];fp.paragraph_format.tab_stops.add_tab_stop(Pt(CW),WD_TAB_ALIGNMENT.RIGHT)
fp.add_run('Demo academic report • PEB structural engineering training\t');field(fp,'PAGE',1);fp.add_run(' / ');field(fp,'NUMPAGES',24)
for r in fp.runs:r.font.name='Times New Roman';r.font.size=Pt(8)
doc.settings.element.append(xml('updateFields',val='true'))

for n,pg in enumerate(pages,1):
    sc=scales[n-1]
    if n==1:
        p=para('A\nSTRUCTURAL ENGINEERING\nTRAINING REPORT',style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;bookmark(p,1)
        p.paragraph_format.space_before=Pt(20);p.paragraph_format.line_spacing=Pt(29)
        for text,size in [('ON',13),('ANALYSIS, DESIGN AND DETAILING OF A\nPRE-ENGINEERED STEEL BUILDING',16),('Proposed PEB Industrial Building, Chennai',12)]:
            p=para(text,size);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(size+5)
        picture('model',185)
        table(['Report particulars','Demo details'],[['Student','Demo Student | Roll No. DEMO-001'],['Institution','Demo College of Engineering'],['Training organization','Demo Structural Consultants'],['Supervisor','Demo Training Supervisor'],['Training period','Four weeks — dates to be entered'],['Report issue','21 October 2026 | Revision D1']],[.30,.70])
        p=para('36 m × 72.2 m • Two 18 m roof spans • Four-week study',11);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=para('Demonstration personal details. Prepared from the supplied project records.\nStudy drawings and illustrative calculations are identified in the report.',10);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        continue
    band(pg['title'],n)
    if 'toc' in pg:
        for level,label,target in pg['toc']:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.left_indent=Pt(0 if level==0 else 12 if level==1 else 23)
            p.paragraph_format.space_after=Pt(4 if level==0 else 0);p.paragraph_format.line_spacing=Pt(14.6)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(CW-3),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
            r=p.add_run(label);r.font.size=Pt(10.2 if level==0 else 10.5);r.bold=level==0
            p.add_run('\t');field(p,f'PAGEREF report_page_{target} \\h',target)
        continue
    for b in pg['blocks']:
        k=b['kind']
        if k=='p':
            front=n in [2,5,6];p=para(b['text'],(12 if n==2 else 11.4 if front else 10.8)*sc)
            p.paragraph_format.line_spacing=Pt((20 if n==2 else 16.5 if front else 13.7)*sc);p.paragraph_format.space_after=Pt(10 if front else 6)
        elif k=='h':
            p=para(b['text'],style='Heading 2');p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4)
        elif k=='source':
            p=para('Source basis: '+b['text'],8.2*sc);p.paragraph_format.line_spacing=Pt(10.5*sc);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif k in ['eq','note']:
            p=para(b['text'],(10.2 if k=='eq' else 9.7)*sc);p.paragraph_format.line_spacing=Pt(13*sc);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent=Pt(6);p.paragraph_format.right_indent=Pt(6);p._p.get_or_add_pPr().append(xml('shd',fill='F2F4F8' if k=='eq' else 'F7F7F7'))
        elif k=='table':table(b['headers'],b['rows'],b.get('widths'),sc)
        elif k=='fig':
            picture(b['name'],b['height']*sc);p=para(b['caption'],8.6*sc);p.paragraph_format.line_spacing=Pt(11*sc);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        elif k=='reference':
            p=para('<b>'+b['label']+'</b> '+b['text'],9.1*sc);p.paragraph_format.line_spacing=Pt(11.8*sc)
            if b['url']:
                rid=p.part.relate_to(b['url'],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
                link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=xml('rPr');pr.append(xml('color',val='4668AC'));pr.append(xml('sz',val=18));r.append(pr);txt=xml('t');txt.text=' [Source]';r.append(txt);link.append(r);p._p.append(link)

doc.core_properties.title='Structural Engineering Training Report — PEB Design (Updated)'
doc.core_properties.author='Demo Student';doc.core_properties.subject='Updated report with abstract, contents and numbered chapters'
updated=OUT/'PEB_Structural_Engineering_Training_Report_Updated.docx';doc.save(updated)
canonical=OUT/'PEB_Structural_Engineering_Training_Report.docx';backup=OUT/'archive'/'PEB_Training_Report_Original.docx'
if canonical.exists() and not backup.exists():shutil.copy2(canonical,backup)
shutil.copy2(updated,canonical)
with zipfile.ZipFile(updated) as z:
    assert z.testzip() is None
    text=z.read('word/document.xml').decode()
    assert text.count('w:bookmarkStart')==24
    assert text.count('PAGEREF report_page_')==63
check=Document(updated)
assert any('1.8 Documentation' in p.text for p in check.paragraphs)
assert len(check.inline_shapes)==7
print('Created:',updated)
print('Verified editable text, tables, 7 figures, 24 page bookmarks and 63 contents references.')
print('Word pagination may vary by application; contents use updateable page-reference fields.')
